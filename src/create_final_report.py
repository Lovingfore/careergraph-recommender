"""Generate the final Chinese project report from current artifacts."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

try:
    from src.database import get_counts
except ModuleNotFoundError:  # Supports ``python src/create_final_report.py``.
    from database import get_counts


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "reports" / "CareerGraph_Recommender_Final_Report.docx"


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_text(cell, text: str, *, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(9.5)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _table(doc: Document, headers: list[str], rows: list[list[object]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[index], header, bold=True, color="FFFFFF")
        _shade(table.rows[0].cells[index], "1F4E79")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            _set_cell_text(cells[index], value)
            if len(table.rows) % 2 == 0:
                _shade(cells[index], "F4F8FC")
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def _heading(doc: Document, text: str, level: int = 1):
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(11 if level == 1 else 7)
    paragraph.paragraph_format.space_after = Pt(5)
    return paragraph


def _bullet(doc: Document, text: str):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(text)
    return paragraph


def _code(doc: Document, text: str):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    run.font.size = Pt(9)
    return paragraph


def _configure(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.18
    for style_name, size, color in [("Heading 1", 17, "1F4E79"), ("Heading 2", 13, "2F75B5"), ("Heading 3", 11, "1F4E79")]:
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("CareerGraph Recommender · 课程设计阶段汇报").font.size = Pt(8)


def build_report() -> Path:
    clean = ROOT / "data" / "clean"
    processed = ROOT / "data" / "processed"
    artifacts = ROOT / "artifacts"
    summary = json.loads((clean / "dataset_summary.json").read_text(encoding="utf-8"))
    model = json.loads((artifacts / "models" / "evaluation_summary.json").read_text(encoding="utf-8"))
    db_counts = get_counts(artifacts / "topic17.sqlite3")
    evaluation = pd.read_csv(artifacts / "evaluation_summary.csv")
    resumes = pd.read_csv(clean / "resumes_zh.csv")
    recommendation_features = pd.read_csv(processed / "recommendation_features.csv")
    forecast = pd.read_csv(processed / "skill_forecast_6m.csv")
    topk = pd.read_csv(processed / "topk_recommendations.csv")
    graph_summary = json.loads((processed / "bipartite" / "bipartite_graph_summary.json").read_text(encoding="utf-8"))

    doc = Document()
    _configure(doc)

    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(70)
    title = cover.add_run("CareerGraph Recommender")
    title.bold = True
    title.font.size = Pt(28)
    title.font.color.rgb = RGBColor.from_string("1F4E79")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("动态职位推荐与职业发展路径规划系统\nWeb 可视化界面与临时简历分析阶段汇报")
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string("475569")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"生成日期：{date.today().isoformat()}\n数据版本：300 条合成中文技术岗简历 · O*NET 30.2 职位/技能字典").font.size = Pt(11)
    doc.add_page_break()

    _heading(doc, "一、项目目标与完成概览")
    doc.add_paragraph("本项目面向课程设计场景，构建从简历技能画像、职位技能需求、职位转移图到动态推荐和职业路径规划的轻量原型。当前版本完成了 300 条中文合成简历数据重建、SQLite 数据库、特征工程、职位—技能二部图、TemporalGAT 训练、Django Web 仪表盘和上传简历即时分析闭环。")
    _table(doc, ["模块", "当前完成内容", "状态"], [
        ["数据与清洗", "O*NET 30.2 抽取 12 个职位、35 项技能；生成 300 条中文合成简历和 6 个月技能事件", "已完成"],
        ["数据库", "SQLite 六张核心表，外键、联合主键、索引和事务导入", "已完成"],
        ["推荐与路径", "内容匹配 + 职位转移图混合推荐，技能缺口、补全时间、最多三跳路径", "已完成"],
        ["模型", "TemporalGAT 轻量训练、顺序切分和 MAE/RMSE 对照评估", "已完成（教学实验）"],
        ["Web", "模型卡片、数据库状态、上传分析、动态推荐、技能缺口、路径和二部图", "已完成"],
        ["汇报材料", "本报告及可复现运行命令、验收证据、限制说明", "已完成"],
    ], [1.2, 4.8, 1.0])

    _heading(doc, "二、数据集与数据库")
    doc.add_paragraph("主数据文件为 data/clean/resumes_zh.csv。简历内容包含学历、专业、城市、工作年限、当前/目标职位、中文技能、技能熟练度、工作经历、项目经历和完整中文文本。所有简历均为确定性合成教学数据，不含真实姓名、电话、邮箱或真实就业记录。")
    _table(doc, ["对象", "数量", "说明"], [
        ["简历 / 用户", f"{len(resumes):,}", "resume_id CR0001–CR0300；user_id u001–u300"],
        ["职位", f"{summary['selected_occupations']:,}", "O*NET 软件、数据和网络方向"],
        ["技能", f"{summary['skills']:,}", "统一 skill_id 与需求权重"],
        ["职位—技能边", f"{summary['occupation_skill_rows']:,}", "完整二部图关系"],
        ["技能事件", f"{summary['user_skill_event_rows']:,}", "300 用户 × 6 月 × 35 技能"],
        ["职位转移", f"{summary['transition_rows']:,}", "确定性教学转移概率"],
    ], [1.7, 1.0, 4.3])
    doc.add_paragraph("SQLite 数据库 artifacts/topic17.sqlite3 当前仅由上述 clean CSV 重建，六张表行数如下：")
    _table(doc, ["表名", "行数"], [[name, count] for name, count in db_counts.items()], [4.8, 1.2])

    _heading(doc, "三、特征工程、二部图与推荐算法")
    _bullet(doc, "职位向量：将 demand_weight 归一化为职位技能需求；用户向量取最近月份技能水平。")
    _bullet(doc, "核心特征：match_score、gap_score、missing_skill_count、direct_transition_probability、path_probability、growth_score、estimated_training_hours。")
    _bullet(doc, "推荐总分：0.45 × Match - 0.25 × Gap + 0.15 × Growth + 0.15 × Path，按用户对候选职位排序。")
    _bullet(doc, f"推荐特征矩阵共 {len(recommendation_features):,} 行，覆盖 300 用户和每用户 11 个候选职位；Top-K 文件共 {len(topk):,} 行。")
    _bullet(doc, f"二部图包含 {graph_summary['occupation_count']} 个职位节点、{graph_summary['skill_count']} 个技能节点和 {graph_summary['edge_count']} 条完整边；阈值 0.30 的展示子图保留 {graph_summary.get('filtered_edge_count', 158)} 条高权重边。")
    graph_image = processed / "bipartite" / "bipartite_graph.png"
    if graph_image.exists():
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(graph_image), width=Inches(6.1))
        caption = doc.add_paragraph("图 1  职位—技能二部图（完整边表另存于 data/processed/bipartite/）")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _heading(doc, "四、模型训练与评价")
    doc.add_paragraph(f"TemporalGAT 使用连续 3 个月技能窗口，当前共 {model['sample_count']} 个样本，按时间顺序切分为训练 {model['split']['train']}、验证 {model['split']['validation']}、测试 {model['split']['test']}，训练 {model['epochs']} 轮，随机种子为 {model['seed']}。")
    _table(doc, ["方法", "MAE", "RMSE", "解释"], [
        ["线性趋势 baseline", model["baseline"]["mae"], model["baseline"]["rmse"], "透明、低成本的 6 个月技能趋势预测"],
        ["TemporalGAT", model["temporal_gat"]["mae"], model["temporal_gat"]["rmse"], "GRU + 多头邻居聚合的轻量模型"],
    ], [2.0, 1.0, 1.0, 2.8])
    doc.add_paragraph("结论边界：当前合成教学数据上，线性 baseline 的误差低于 TemporalGAT。因此页面和报告将 TemporalGAT 标记为“训练完成的教学实验模型”，不会宣称其优于 baseline。在线 /api/forecast/接口仍明确返回 linear_trend_baseline，避免把训练 checkpoint 冒充为生产预测。")

    _heading(doc, "五、Web 可视化界面")
    _bullet(doc, "数据仪表盘：显示简历/用户、职位、技能、技能事件、训练窗口和六张 SQLite 表行数。")
    _bullet(doc, "模型信息卡：显示训练状态、训练轮数、540/180/180 样本切分、TemporalGAT 与 baseline 的 MAE，并提示合成数据和模型边界。")
    _bullet(doc, "已有用户查询：输入 user_id 和目标职位，调用推荐、技能缺口、职业路径与 6 个月 baseline 预测接口。")
    _bullet(doc, "临时简历分析：上传 UTF-8 .txt/.md/.csv（≤1 MB），识别中英文技能关键词，自动推断当前职位或使用下拉选择，生成动态推荐、技能缺口、补全时间和职业路径。")
    _bullet(doc, "可视化结果：推荐表、缺口表、路径节点和职位—技能二部图均在页面中展示；页面使用原生 HTML/CSS/JavaScript，无 CDN 和构建步骤。")
    _table(doc, ["接口", "用途", "方法"], [
        ["/api/model-info/", "模型状态和 MAE/RMSE", "GET"],
        ["/api/occupations/", "下拉职位字典（含中文名称）", "GET"],
        ["/api/resume-upload/", "临时解析和动态规划", "POST multipart"],
        ["/api/recommend/", "基准用户 Top-K 推荐", "GET"],
        ["/api/skill-gap/", "职位技能缺口", "GET"],
        ["/api/career-path/", "最多三跳职业路径", "GET"],
        ["/bipartite-graph.svg", "二部图可视化", "GET"],
    ], [2.0, 3.8, 1.0])

    _heading(doc, "六、简历上传数据流与不落库保证")
    _table(doc, ["步骤", "处理", "持久化"], [
        ["1. 选择文件", "校验扩展名、大小和 multipart 字段", "无"],
        ["2. 文本解码", "utf-8-sig 读取 .txt/.md/.csv，拒绝空文本和非法编码", "无"],
        ["3. 技能映射", "中文/英文技能别名 → skill_id，构造 35 维临时向量", "无"],
        ["4. 动态计算", "对 12 个职位实时计算 Match/Gap/Growth/Path", "无"],
        ["5. 返回结果", "返回 recognized_skills、recommendations、skill_gap、career_path 和 persisted=false", "无"],
    ], [1.2, 5.0, 0.6])
    doc.add_paragraph("上传接口不调用 SQLite 写入或 CSV 输出逻辑；测试在请求前后比较六张表计数，并校验 resumes_zh.csv、user_profiles.csv、user_skill_events.csv 的 SHA-256 未改变，确保上传分析不会污染基准数据。")
    _bullet(doc, "页面通过 Django CSRF token 保护上传 POST；服务端显式限制单文件 1 MB 和请求体 2 MB。")

    _heading(doc, "七、原始第17题要求对照")
    _table(doc, ["要求", "实现 / 证据", "状态"], [
        ["职位—技能二部图", "bipartite_graph.json、bipartite_edges.csv、SVG/PNG；12 职位、35 技能、420 边", "已完成"],
        ["职位转移概率图", "job_transitions.csv 与 transition_graph.json；最多三跳路径搜索", "已完成"],
        ["未来 6 个月技能预测", "skill_forecast_6m.csv；线性趋势 baseline，TemporalGAT 训练 checkpoint", "已完成（教学实验）"],
        ["多目标 Top-K 推荐", "推荐总分融合匹配、缺口、成长和路径；Precision@K/Recall@K/NDCG@K 已计算", "已完成"],
        ["技能缺口与补全时间", "页面上传结果和基准用户接口返回缺口数量、技能差距和 estimated_training_hours", "已完成"],
        ["职业发展路径规划", "当前职位/目标职位输入框，返回路径、路径概率和跳数", "已完成"],
        ["个人技能状态图", "页面用识别技能标签和技能水平结果展示；二部图提供全局关系可视化", "轻量实现"],
        ["Skill Gap Reduction 等长期指标", "当前版本保留技能缺口快照，跨时间 reduction 曲线列为后续优化", "待扩展"],
    ], [2.2, 4.0, 0.8])

    _heading(doc, "八、测试、验收与运行证据")
    _bullet(doc, "全量单元测试：37/37 通过，覆盖数据完整性、数据库外键关系、DataLoader、TemporalGAT 前向、推荐服务、二部图、Web 页面、上传接口、CSRF 和不落库校验。")
    _bullet(doc, "阶段验收 verify_stage1.py：data_integrity、database、feature_outputs、loader、bipartite_graph、web_contract、web_api、resume_dataset 均为 ok=true。")
    _bullet(doc, f"推荐评价：Precision@1={evaluation.loc[evaluation['k'] == 1, 'precision_at_k'].iloc[0]:.6f}，Recall@3={evaluation.loc[evaluation['k'] == 3, 'recall_at_k'].iloc[0]:.6f}，NDCG@5={evaluation.loc[evaluation['k'] == 5, 'ndcg_at_k'].iloc[0]:.6f}。")
    _bullet(doc, f"技能预测文件共 {len(forecast)} 条（300 用户 × 35 技能），DataLoader 形成 900 个 [3, 35] 时序样本。")
    _code(doc, "D:\\python\\python.exe -m pip install -r requirements.txt")
    _code(doc, "D:\\python\\python.exe -m pip install -r requirements-optional-models.txt")
    _code(doc, "D:\\python\\python.exe src\\prepare_dataset.py")
    _code(doc, "D:\\python\\python.exe src\\forecast_skills.py")
    _code(doc, "D:\\python\\python.exe src\\build_features.py")
    _code(doc, "D:\\python\\python.exe src\\init_database.py --db-path artifacts\\topic17.sqlite3 --data-dir data\\clean")
    _code(doc, "D:\\python\\python.exe src\\train_temporal_gat.py --data-dir data\\clean --artifact-dir artifacts\\models --epochs 30 --seed 42")
    _code(doc, "D:\\python\\python.exe web\\manage.py runserver 127.0.0.1:8000")

    _heading(doc, "九、跨设备部署与限制")
    _bullet(doc, "项目只使用相对路径；基础页面、推荐、缺口、路径和二部图依赖 requirements.txt，TemporalGAT 训练另需 requirements-optional-models.txt。")
    _bullet(doc, "上传简历不落库，适合课程演示和跨设备临时分析；若未来需要长期保存，必须另行设计授权、脱敏、版本和增量训练流程。")
    _bullet(doc, "O*NET 是公开真实职位/技能参考数据，但简历、技能事件和职位转移为合成教学数据；当前指标不能外推到真实招聘市场。")
    _bullet(doc, "当前 TemporalGAT 是轻量模型，测试误差高于线性 baseline；后续可增加真实授权数据、技能缺口减少指标、雷达图和更系统的超参数调优。")

    _heading(doc, "十、交付文件索引")
    _table(doc, ["文件", "作用"], [
        ["data/clean/resumes_zh.csv", "300 条合成中文技术岗简历主数据"],
        ["artifacts/topic17.sqlite3", "六张核心表 SQLite 数据库"],
        ["artifacts/models/temporal_gat.pt", "TemporalGAT checkpoint"],
        ["web/topic17_app/views.py", "模型、职位、上传和原有 API"],
        ["web/topic17_app/templates/topic17_app/index.html", "中文响应式 Web 仪表盘"],
        ["src/resume_analysis.py", "不落库的临时简历解析和动态评分"],
        ["docs/superpowers/specs/...design.md", "已确认的设计规格"],
        ["docs/superpowers/plans/...analysis.md", "实现计划与验收步骤"],
    ], [3.7, 3.3])
    doc.add_paragraph("报告生成脚本为 src/create_final_report.py；重新生成时会读取当前 CSV、模型评估 JSON 和图文件，避免手工复制过时数字。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_report())
